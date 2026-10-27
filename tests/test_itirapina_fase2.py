"""
tests/test_itirapina_fase2.py — Testes de aceitação da Fase 2 do trecho
Itirapina–Pederneiras (PROMPT_VSCode_Sentinel_Itirapina_OnePage_v1, seção 9).

Fixtures reais em tests/fixtures/itirapina/ (fora do git); sem a pasta, os
testes que dependem delas são pulados.

Diferença documentada:
  - Top 20 do plano = 1.657 dormentes (e não 1.643): 78/300818-302202tg
    aparece 2× na LP da Engenharia (158 + 14); o estudo rankeou só a 1ª
    linha. O travamento total (4.645) só fecha somando as duas.
  - Vagão: 153 corretivos · 1.680 perímetro urbano · 869 com lat/long são da
    planilha inteira (Sub 78 toda); no trecho são 105 · 1.557 · 664.
"""

import sys
from datetime import date
from pathlib import Path

import pandas as pd
import pytest

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))
FIX = RAIZ / "tests" / "fixtures" / "itirapina"
com_fixtures = pytest.mark.skipif(not FIX.exists(), reason="fixtures do trecho ausentes")


def _csv(nome):
    return pd.read_csv(FIX / "dados" / nome, sep=";", encoding="utf-8-sig")


@pytest.fixture(scope="module")
def rel():
    from core.parser_relatorio_velocidade import ler_relatorio_velocidade
    from core.analise_velocidade import normalizar_relatorio
    rota, exp, df = ler_relatorio_velocidade(FIX / "JUNDIAI-TRIAGEM PAULISTA.pdf")
    return normalizar_relatorio(df.assign(rota=rota))


@pytest.fixture(scope="module")
def defeitos_vagao():
    from core.parser_defeitos_via import ler_defeitos_via
    return ler_defeitos_via(FIX / "Vgao Via-Defeitos Pederneiras 2610.xlsx")[0]


@pytest.fixture(scope="module")
def plano():
    from core.plano_itirapina import ler_plano_engenharia
    p, _ = ler_plano_engenharia(FIX / "Prospecção Itirapina Engenharia.xlsx",
                                _csv("pontos_criticos_C_T.csv"), _csv("top10_curvas_tangentes.csv"))
    return p


# region ── F0 ───────────────────────────────────────────────────────────────

def test_sb_pelo_km():
    from core.analise_trecho import sb_pelo_km, SB_FAIXAS
    assert sb_pelo_km(206.0) == "ZWF-ZBO"
    assert sb_pelo_km(290.0) == "ZWM-ZWY"
    assert sb_pelo_km(300.0) == "ZWY-ZPD"
    assert sb_pelo_km(303.0) == "ZPD-ZFJ"
    assert sb_pelo_km(303.8) == "ZPD-ZFJ"
    assert sb_pelo_km(174.8) == "ZIQ-ZQY"
    assert sb_pelo_km(150.0) is None and sb_pelo_km(None) is None
    assert len(SB_FAIXAS) == 16


def test_preencher_sb_so_onde_vazia():
    from core.analise_trecho import preencher_sb
    d = pd.DataFrame({"km_inicio": [206.0, 300.0, 290.0], "km_fim": [206.2, 300.4, 290.1],
                      "sb": [None, "None", "XPTO"]})
    f = preencher_sb(d, "km_inicio", "km_fim")
    assert f["sb"].tolist() == ["ZWF-ZBO", "ZWY-ZPD", "XPTO"]


@com_fixtures
def test_matriz_sb_16_com_rev1_sem_sb():
    """Rev 1 não tem SB: os indicadores por SB vêm pelo km → 16 SBs."""
    from core.parser_prospeccao_campo import ler_prospeccao_campo
    from core.parser_prospeccao import enriquecer_prospeccao
    from core.indicadores_trecho import calcular
    abas, _ = ler_prospeccao_campo(FIX / "Prospecção - Sub 78 Rev (1).xlsx")
    dorm = enriquecer_prospeccao(abas["dormentes"])
    foto = calcular({"dormentes": dorm}, date(2026, 10, 26))
    sb = foto[(foto["escopo"] == "sb") & (foto["indicador"] == "dormentes_prospectados")]
    assert sb["chave"].nunique() == 16
    assert (sb["valor"] > 0).sum() == 16


@com_fixtures
def test_restricoes_unica_fonte(rel):
    from core.analise_trecho import restricoes_trecho, rotulo_restricoes
    rA = restricoes_trecho(rel)
    assert rotulo_restricoes(rA) == "30 (14 via · 16 ANTT)"


@com_fixtures
def test_vagao_recarregado(defeitos_vagao, rel):
    from core.analise_trecho import vagao_sem_restricao
    v = vagao_sem_restricao(defeitos_vagao, rel, viagem=180)
    assert int((~v["coberto"]).sum()) == 52
    # O prompt cita 153 corretivos · 1.680 PU · 869 lat/long: são da planilha
    # inteira (7.928 linhas, Sub 78 toda). No trecho (km ≤ 303,8, 6.360):
    assert int(defeitos_vagao[["latitude", "longitude"]].notna().all(axis=1).sum()) == 664
    assert int((defeitos_vagao["classificacao"] == "Corretivo").sum()) == 105
    assert int((defeitos_vagao["classificacao"] == "Crítico").sum()) == 2
    assert int((defeitos_vagao["perimetro_urbano"] == True).sum()) == 1557  # noqa: E712
    assert int((defeitos_vagao["acidente"] == True).sum()) == 151  # noqa: E712

# endregion


# region ── F1 ───────────────────────────────────────────────────────────────

@com_fixtures
def test_plano_totais(plano):
    from core.plano_itirapina import resumo_plano
    r = resumo_plano(plano)
    assert r["travamento"] == 4645
    assert r["top20_locais"] == 20
    assert r["top20_travamento"] == 1643 + 14      # ver docstring
    assert abs(r["top20_km"] - 22.3) < 0.06
    assert r["pontos"] == ["C1", "C2", "C3", "T1", "T2", "T3", "T4"]
    pts = _csv("pontos_criticos_C_T.csv").set_index("ponto")["travamento"].to_dict()
    assert plano.groupby("ponto_critico")["dorm_travamento"].sum().to_dict() == pts


@com_fixtures
def test_execucao_sobe_avanco_t1(plano):
    from core.plano_itirapina import avanco_travamento
    antes = avanco_travamento(plano, None)
    t1 = plano[(plano["ponto_critico"] == "T1") & (plano["dorm_travamento"] > 0)].iloc[0]
    exe = pd.DataFrame([{"data": "2027-01-10", "atividade": "troca_dormente",
                         "km_inicio": (t1.km_inicio + t1.km_fim) / 2, "linha": "LP",
                         "equipamento_siv": None, "quantidade": 100, "unidade": "dorm"}])
    depois = avanco_travamento(plano, exe)
    assert antes["executado"] == 0 and depois["executado"] == min(100, int(t1.dorm_travamento))
    assert depois["pct"] > antes["pct"]
    p = depois["pontos"].set_index("ponto_critico")
    assert p.loc["T1", "pct"] > 0 and p.loc["T1", "situacao"] in ("em curso", "concluído")


def test_modelo_execucao_ida_e_volta(tmp_path):
    from core.parser_execucao import gerar_modelo, ler_execucao
    arq = tmp_path / "modelo.xlsx"
    arq.write_bytes(gerar_modelo())
    df, avisos = ler_execucao(arq)
    assert avisos == []
    assert df["atividade"].tolist() == ["troca_dormente", "solda"]
    assert df["unidade"].tolist() == ["dorm", "und"]
    assert df["linha"].tolist() == ["LP", "LP"]
    assert str(df["data"].iloc[0]) == "2027-01-15"


def test_execucao_km_e_linhas_invalidas(tmp_path):
    from core.parser_execucao import ler_execucao
    arq = tmp_path / "e.xlsx"
    pd.DataFrame({"Data": ["10/02/2027", "11/02/2027", None, "12/02/2027"],
                  "Atividade": ["troca_dormente", "Socaria", "Solda", "xpto"],
                  "KM início": ["285+100", 285100, 286, 287],
                  "Quantidade": [10, 0.5, 2, 1]}).to_excel(arq, index=False)
    df, avisos = ler_execucao(arq)
    assert df["km_inicio"].tolist() == [285.1, 285.1]
    assert df["unidade"].tolist() == ["dorm", "km"]
    assert len(avisos) == 1 and "2 linha(s) ignorada(s)" in avisos[0]

# endregion


# region ── F2 ───────────────────────────────────────────────────────────────

@com_fixtures
def test_top10_reproduz_estudo(plano):
    from core.plano_itirapina import top_n
    t = top_n(plano, 10)
    csv = _csv("top10_curvas_tangentes.csv")
    for tipo in ("curva", "tangente"):
        assert t[t["tipo"] == tipo]["equipamento_siv"].tolist() == \
            csv[csv["tipo"] == tipo].sort_values("rank")["equipamento_SIV"].tolist()


@com_fixtures
def test_pontos_criticos_algoritmo(plano):
    """Reproduz C1, T1 e T2 (km ± 0,1). C2, C3, T3 e T4 do estudo tiveram
    ajuste manual (ex.: a janela do km 251,5 tem 932 inservíveis > C3 = 919)."""
    from core.plano_itirapina import pontos_criticos
    calc = pontos_criticos(plano)
    assert calc["ponto"].tolist() == ["C1", "C2", "C3", "T1", "T2", "T3", "T4"]
    oficial = _csv("pontos_criticos_C_T.csv").set_index("ponto")
    achados = {(round(r.km_inicio, 1), round(r.km_fim, 1)) for r in calc.itertuples()}
    for p in ("C1", "T1", "T2"):
        o = oficial.loc[p]
        assert any(abs(a - o.km_ini) <= 0.1 and abs(b - o.km_fim) <= 0.1 for a, b in achados), p


@com_fixtures
def test_quadro_equipamentos(plano):
    from core.plano_itirapina import quadro_equipamentos
    q = quadro_equipamentos(plano)
    assert len(q) == len(plano) and q["dorm_travamento"].sum() == 4645
    assert (q["pct_exec"] == 0).all()
    assert q["taxa_pct"].between(0, 100).all()

# endregion


# region ── F3 ───────────────────────────────────────────────────────────────

def test_lod_e_janela():
    from modules.itirapina.unifilar import lod_da_janela, motivo_curto
    from modules.itirapina.navegacao import _limitar
    assert lod_da_janela(174.8, 303.8) == "agregado"
    assert lod_da_janela(280, 300) == "medio"
    assert lod_da_janela(284, 286) == "detalhe"
    assert _limitar(173.0, 175.0) == (174.8, 176.8)        # não sai do trecho
    assert _limitar(303.0, 305.0) == (301.8, 303.8)
    assert motivo_curto("ANTT - PERIMETRO URBANO") == "PERIMETRO URBANO"

# endregion


# region ── F4 ───────────────────────────────────────────────────────────────

def _rel_modificado(rel):
    """Cópia do relatório sem 2 restrições de via da LP e com 1 velocidade mudada."""
    from core.analise_trecho import TRECHO_KM, GRUPO_ANTT
    r = rel.copy()
    m = ((r["tipo"] == "restricao_a") & (r["sub"] == "78") & (r["categoria"] != GRUPO_ANTT)
         & (r["linha"] == "LP") & (r[["km_de", "km_ate"]].min(axis=1) <= TRECHO_KM[1]))
    idx = r[m].sort_values("km_de").index
    r = r.drop(index=idx[:2])
    r.loc[idx[2], "vel"] = r.loc[idx[2], "vel"] + 5
    return r


@com_fixtures
def test_eventos_de_restricao(rel):
    from core.eventos_restricao import comparar_restricoes, resumo_eventos
    ev = comparar_restricoes(rel, _rel_modificado(rel))
    assert (ev["evento"] == "retirada").sum() == 2
    assert (ev["evento"] == "alterada").sum() == 1
    assert (ev["evento"] == "nova").sum() == 0
    assert resumo_eventos(ev)["via_retiradas"] == 2
    assert set(ev["grupo"]) == {"via"}
    assert comparar_restricoes(rel, rel).empty


@com_fixtures
def test_minutos_ganhos(rel):
    from core.indicadores_trecho import calcular, derivados
    from core.eventos_restricao import comparar_restricoes
    a = calcular({"rel": rel}, date(2026, 10, 26))
    novo = _rel_modificado(rel)
    b = calcular({"rel": novo}, date(2027, 3, 1))
    d = derivados(a, b, comparar_restricoes(rel, novo))
    assert d["minutos_ganhos"] > 0
    assert d["km_liberados"] > 0
    assert d["via_retiradas"] == 2


def test_motivo_legivel_e_grupo():
    from core.eventos_restricao import motivo_legivel, grupo_evento, codigo_motivo
    assert motivo_legivel("509 VIA DESNIVELADA") == "via desnivelada"
    assert motivo_legivel("502 ANTT - PN COM SINALIZAÇÃO INSUFICIENTE") == "PN sem sinalização"
    assert grupo_evento(codigo_motivo("503 ANTT - PERIMETRO URBANO")) == "ANTT/terceiros"
    assert grupo_evento(codigo_motivo("204 TORÇÃO")) == "via"

# endregion


# region ── F5/F6 ────────────────────────────────────────────────────────────

@pytest.fixture(scope="module")
def onepages():
    sys.path.insert(0, str(RAIZ / "tests"))
    from _onepage_fix import insumos
    from core.onepage import montar_onepage
    return {m: montar_onepage(**insumos(m == "evolucao")) for m in ("partida", "evolucao")}


@com_fixtures
def test_onepage_modos(onepages):
    p, e = onepages["partida"], onepages["evolucao"]
    assert p["modo"] == "partida" and e["modo"] == "evolucao"
    assert all(c["antes"] is None and c["delta"] is None for c in p["cards"])      # sem setas
    assert (p["restricoes"]["status"] == "ativa").sum() == 30
    assert (e["restricoes"]["status"] == "retirada").sum() == 2
    assert any("pontos críticos" in f for f in p["frases"])                       # frases do plano


@com_fixtures
@pytest.mark.parametrize("modo", ["partida", "evolucao"])
def test_onepage_exporta(onepages, modo):
    import io
    from PIL import Image
    from core.onepage_render import (render_onepage, para_png, para_pdf, para_pptx, textos_fora_do_quadro,
                                     nome_arquivo)
    o = onepages[modo]
    assert textos_fora_do_quadro(render_onepage(o)) == []
    png = para_png(o)
    im = Image.open(io.BytesIO(png))
    assert im.size == (2338, 1654) and getattr(im, "n_frames", 1) == 1
    assert para_pdf(o).startswith(b"%PDF")
    from pptx import Presentation
    prs = Presentation(io.BytesIO(para_pptx(o, png)))
    assert len(prs.slides) == 1 and prs.slides[0].notes_slide.notes_text_frame.text
    assert nome_arquivo(o, "png").startswith("OnePage_Itirapina_Pederneiras_")


def test_frases_nunca_sobre_piora():
    from core.onepage_textos import frases_evolucao
    piorou = {"via_retiradas": 0, "minutos_ganhos": -4, "horas_mes": -20, "travamento_pct": 0,
              "defeitos_eliminados": 0, "km_liberados": -1}
    base = {"restricoes_via_mz": 14, "vma_atual": 37, "taxa_mz": 25.8, "taxa_atual": 28.0,
            "vagao_mz": 52, "vagao_atual": 60}
    assert frases_evolucao(piorou, base) == []
    # ganho insignificante também não vira frase
    assert frases_evolucao({"minutos_ganhos": 0.4, "travamento_pct": 2}, {"taxa_mz": 25.8, "taxa_atual": 25.5}) == []
    ok = frases_evolucao({"via_retiradas": 3, "km_liberados": 2.0, "minutos_ganhos": 9, "horas_mes": 87},
                         {"restricoes_via_mz": 14, "vma_atual": 40})
    assert len(ok) == 2 and "87 h de trem por mês" in ok[1]

# endregion
