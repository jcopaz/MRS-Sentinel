# =============================================================================
# core/onepage_render.py — One Page em matplotlib (exportação PNG/PDF/PPTX)
#
# Sem kaleido/Chrome (instáveis no Streamlit Cloud): a página é redesenhada em
# matplotlib a partir do MESMO dict de core/onepage.montar_onepage, com as
# cores de core/cores_itirapina. A4 paisagem (11,69 × 8,27 in), 200 dpi,
# DejaVu Sans. Layout fixo em frações da figura (um único quadro).
#   render_onepage(o) -> Figure · para_png / para_pdf / para_pptx -> bytes
#   nome_arquivo(o, ext) -> OnePage_Itirapina_Pederneiras_aaaa-mm-dd.ext
# =============================================================================

import io
import textwrap

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib.patches import FancyBboxPatch, Rectangle  # noqa: E402

from core import cores_itirapina as cores  # noqa: E402
from core.analise_trecho import TRECHO_KM  # noqa: E402
from core.onepage_textos import _n, sem_markdown  # noqa: E402

A4 = (11.69, 8.27)
DPI = 200
COR_CARD = {"bom": "#16a34a", "ruim": "#dc2626", "neutro": "#0f172a"}
plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 7, "axes.spines.top": False,
                     "axes.spines.right": False})


def _titulo_bloco(fig, x, y, txt):
    fig.text(x, y, txt, fontsize=7.5, fontweight="bold", color=cores.AZUL, va="bottom")


def _cabecalho(fig, o):
    fig.patches.append(Rectangle((0.015, 0.925), 0.97, 0.06, transform=fig.transFigure,
                                 facecolor=cores.AZUL, edgecolor="none"))
    fig.text(0.025, 0.958, f"{o['titulo']}  ({o['subtitulo']})", color="white", fontsize=12, fontweight="bold",
             va="center")
    fig.text(0.025, 0.936, o["linha_datas"], color="#e2e8f0", fontsize=8, va="center")
    fig.text(0.975, 0.947, "MRS · Sentinel", color="white", fontsize=9, ha="right", va="center")


def _cards(fig, o):
    cs = o["cards"]
    n = len(cs)
    larg = (0.97 - 0.01 * (n - 1)) / n
    for i, c in enumerate(cs):
        x = 0.015 + i * (larg + 0.01)
        fig.patches.append(FancyBboxPatch((x, 0.79), larg, 0.12, boxstyle="round,pad=0.003",
                                          transform=fig.transFigure, facecolor="white",
                                          edgecolor=cores.AZUL if c["destaque"] else "#cbd5e1",
                                          linewidth=1.4 if c["destaque"] else 0.8))
        fig.text(x + 0.008, 0.892, c["titulo"], fontsize=6.5, fontweight="bold", color="#475569", va="top")
        val = (f"{c['antes']} → " if c["antes"] else "") + c["valor"]
        fig.text(x + 0.008, 0.848, val, fontsize=12 if len(val) < 16 else 10, fontweight="bold",
                 color=COR_CARD[c["cor"]], va="center")
        if c["delta"]:
            fig.text(x + larg - 0.006, 0.822, c["delta"], fontsize=7.5, color=COR_CARD[c["cor"]], ha="right")
        fig.text(x + 0.008, 0.80, textwrap.shorten(c["rodape"], 46, placeholder="…"), fontsize=5.8,
                 color="#64748b", va="bottom")


def _esquematico(fig, o):
    _titulo_bloco(fig, 0.015, 0.755, "① ESQUEMÁTICO DO TRECHO — cada restrição no km real")
    ax = fig.add_axes([0.03, 0.60, 0.94, 0.155])
    ax.set_xlim(TRECHO_KM[0] - 1, TRECHO_KM[1] + 1)
    ax.set_ylim(-0.6, 1.62)
    ax.axhline(0, color="#334155", lw=3, zorder=1)
    e = o["estacoes"]
    if len(e):
        e = e[(e["km_fim"] >= TRECHO_KM[0] - 6) & (e["km_inicio"] <= TRECHO_KM[1])]
        for r in e.itertuples():
            k = min(max((r.km_inicio + r.km_fim) / 2, TRECHO_KM[0]), TRECHO_KM[1])
            ax.plot([k, k], [-0.13, 0.13], color="#334155", lw=1.5)
            ax.text(k, -0.22, r.codigo, ha="center", va="top", fontsize=6, color="#334155")
    r = o["restricoes"]
    if len(r):
        # ~12 km de largura por rótulo nesta escala; 5 "trilhos" de rótulo
        niveis, meia = 5, 6.5
        fim = [-1e9] * niveis
        for x in r.sort_values("a").itertuples():
            meio = (x.a + x.b) / 2
            t = next((i for i, f in enumerate(fim) if meio - meia > f), None)
            if t is None:
                t = min(range(niveis), key=lambda i: fim[i])
            fim[t] = meio + meia
            cor = cores.RETIRADA if x.status == "retirada" else cores.VIA if x.grupo == "via" else cores.ANTT
            ax.add_patch(Rectangle((x.a - 0.25, -0.1), max(x.b - x.a, 0) + 0.5, 0.2, color=cor, zorder=2))
            y = 0.32 + 0.27 * t
            ax.plot([meio, meio], [0.1, y - 0.07], color=cor, lw=0.5)
            txt = f"{x.vel:.0f} · {textwrap.shorten(str(x.motivo), 20, placeholder='…')} · {_n(x.a, 1)}"
            if x.status == "retirada":
                txt = "✔ " + txt
            borda = (TRECHO_KM[1] - TRECHO_KM[0]) * 0.06
            ha = "left" if meio - TRECHO_KM[0] < borda else "right" if TRECHO_KM[1] - meio < borda else "center"
            ax.text(meio, y, txt, ha=ha, va="center", fontsize=4.8, color=cor,
                    bbox=dict(facecolor="white", edgecolor="none", pad=0.4, alpha=0.85))
    ax.set_yticks([])
    ax.spines["left"].set_visible(False)
    ax.tick_params(axis="x", labelsize=5.5)
    fig.text(0.03, 0.553, "■ restrição de via (Manutenção resolve)", color=cores.VIA, fontsize=6)
    fig.text(0.24, 0.553, "■ ANTT/PN/perímetro urbano (terceiros)", color=cores.ANTT, fontsize=6)
    fig.text(0.45, 0.553, "■ retirada", color=cores.RETIRADA, fontsize=6)


def _dormentacao(fig, o):
    _titulo_bloco(fig, 0.015, 0.515, "② DORMENTAÇÃO INSERVÍVEL por equipamento"
                  + (" (fino = marco zero)" if o.get("dorm_ref") is not None else ""))
    ax = fig.add_axes([0.04, 0.345, 0.44, 0.16])
    d = o["dorm"]
    for linha, y in (("LP", 1.0), ("LD", 0.0)):
        g = d[d["linha"] == linha]
        for r in g.itertuples():
            ax.add_patch(Rectangle((r.km_inicio, y - 0.25), abs(r.km_fim - r.km_inicio), 0.5, linewidth=0,
                                   color=cores.cor_por_faixa(r.taxa_pct, cores.LIMIARES_SLIDE, cores.CORES_SLIDE)))
        ref = o.get("dorm_ref")
        if ref is not None and len(ref):
            for r in ref[ref["linha"] == linha].itertuples():
                ax.add_patch(Rectangle((r.km_inicio, y + 0.32), abs(r.km_fim - r.km_inicio), 0.1, linewidth=0,
                                       color=cores.cor_por_faixa(r.taxa_pct, cores.LIMIARES_SLIDE,
                                                                 cores.CORES_SLIDE)))
    p = o["pontos"]
    if len(p):
        for x in p[p["ponto"].astype(str).str.match(r"^[CT]\d$")].itertuples():
            ax.axvspan(x.km_inicio, x.km_fim, color=cores.AZUL, alpha=0.1, lw=0)
            ax.text(x.km_inicio, 1.5, x.ponto, rotation=90, fontsize=5, va="bottom", color=cores.AZUL)
    ax.set_xlim(*TRECHO_KM)
    ax.set_ylim(-0.5, 1.75)
    ax.set_yticks([0, 1], ["LD", "LP"])
    ax.tick_params(labelsize=5.5)
    rot = ["≤10%", "10–20%", "20–30%", "30–40%", ">40%"]
    for i, (c, t) in enumerate(zip(cores.CORES_SLIDE, rot)):
        fig.text(0.04 + i * 0.06, 0.298, "■ " + t, color=c, fontsize=5.5)


def _transit(fig, o):
    _titulo_bloco(fig, 0.52, 0.515, "③ TRANSIT TIME MÊS A MÊS (tracejado = marco zero)")
    ax = fig.add_axes([0.55, 0.345, 0.42, 0.16])
    paleta = {"ZIQ→ZPD": cores.AZUL, "ZPD→ZIQ": cores.LARANJA}
    t = o["transit"]
    for s, g in t.groupby("sentido"):
        g = g.sort_values("mes")
        cor = paleta.get(s, cores.NEUTRO)
        ax.plot(g["mes"], g["mediana"], color=cor, lw=1.6, marker="o", ms=2, label=f"{s} mediana")
        ax.plot(g["mes"], g["media"], color=cor, lw=0.8, ls=":", label=f"{s} média")
        ref = o["transit_ref"].get(s) or {}
        if ref.get("mediana") is not None:
            ax.axhline(ref["mediana"], color=cor, lw=0.7, ls="--")
        u = g.iloc[-1]
        ax.annotate(f"{_n(u['mediana'], 1)} h", (u["mes"], u["mediana"]), fontsize=5.5, color=cor,
                    xytext=(0, 4), textcoords="offset points", ha="center")
    ax.set_ylabel("h", fontsize=6)
    ax.tick_params(labelsize=5.5)
    import matplotlib.dates as mdates
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%m/%y"))
    if len(t):
        ax.legend(fontsize=5, ncol=4, loc="upper center", bbox_to_anchor=(0.5, -0.14), frameon=False)


def _travamento(fig, o):
    _titulo_bloco(fig, 0.015, 0.245, "④ TRAVAMENTO – 7 PONTOS E TOP 10")
    ax = fig.add_axes([0.04, 0.095, 0.30, 0.145])
    p = o["travamento"]
    if len(p):
        p = p.sort_values("ponto_critico", ascending=False)
        ax.barh(p["ponto_critico"], p["pct"] * 100, color=[cores.SITUACAO[s] for s in p["situacao"]], height=0.6)
        for i, r in enumerate(p.itertuples()):
            ax.text(max(r.pct * 100, 0) + 2, i, f"km {_n(r.km_inicio, 1)}–{_n(r.km_fim, 1)} · "
                    f"{_n(r.travamento)} dorm. · {_n(r.pct * 100)}%", va="center", fontsize=5)
    ax.set_xlim(0, 160)
    ax.set_xticks([0, 50, 100])
    ax.tick_params(labelsize=5.5)
    t = o["top"]
    if t is not None and len(t):
        for j, (tipo, rot) in enumerate((("curva", "Top 10 curvas"), ("tangente", "Top 10 tangentes"))):
            y = 0.065 - j * 0.022
            fig.text(0.04, y, rot, fontsize=5.5, color="#475569", va="center")
            g = t[t["tipo"] == tipo].sort_values("rank_calc")
            for k, r in enumerate(g.itertuples()):
                fig.text(0.125 + k * 0.016, y, "●", fontsize=10, ha="center", va="center",
                         color=cores.SITUACAO[r.situacao])
        fig.text(0.30, 0.054, "● concluído", color=cores.SITUACAO["concluído"], fontsize=5.5, va="center")
        fig.text(0.30, 0.040, "● em curso", color=cores.SITUACAO["em curso"], fontsize=5.5, va="center")
        fig.text(0.30, 0.026, "● a iniciar", color="#94a3b8", fontsize=5.5, va="center")


def _feito(fig, o):
    x0 = 0.37
    if o["modo"] == "partida":
        _titulo_bloco(fig, x0, 0.245, "⑤ O PLANO")
        p = o["plano"]
        itens = [f"{p['pontos']} pontos críticos (C1–C3 / T1–T4)", f"{_n(p['travamento'])} dormentes de travamento",
                 f"{p['restricoes_via']} restrições de via a retirar", f"{p['vagao']} pontos do Vagão a tratar"]
    else:
        _titulo_bloco(fig, x0, 0.245, "⑤ O QUE FIZEMOS (desde o marco zero)")
        f = o["feito"]
        itens = [f"{_n(r.total)} {r.unidade} · {r.rotulo}" for r in f.itertuples()] or \
            ["Nenhuma execução registrada"]
    for i, t in enumerate(itens[:6]):
        fig.text(x0 + 0.005, 0.225 - i * 0.022, "• " + t, fontsize=6.5, va="top")
    xf = 0.62
    for i, fr in enumerate(o["frases"][:3]):
        linhas = textwrap.wrap(sem_markdown(fr), 62)
        y = 0.235 - i * 0.06
        fig.patches.append(Rectangle((xf - 0.006, y - 0.012 * len(linhas) - 0.004), 0.003,
                                     0.012 * len(linhas) + 0.004, transform=fig.transFigure,
                                     color=cores.LARANJA))
        fig.text(xf, y, "\n".join(linhas), fontsize=6.6, va="top", color="#0f172a", linespacing=1.25)
    if o["proximos"]:
        fig.text(x0 + 0.005, 0.105, "\n".join(textwrap.wrap("Próximos passos: " + " · ".join(o["proximos"]), 70)),
                 fontsize=5.8, color="#475569", va="top")


def render_onepage(o: dict):
    fig = plt.figure(figsize=A4, dpi=DPI)
    fig.patch.set_facecolor("white")
    _cabecalho(fig, o)
    _cards(fig, o)
    _esquematico(fig, o)
    _dormentacao(fig, o)
    _transit(fig, o)
    _travamento(fig, o)
    _feito(fig, o)
    fig.text(0.015, 0.012, textwrap.shorten(o["fontes"], 230, placeholder="…"), fontsize=5, color="#64748b")
    return fig


def _salvar(o: dict, fmt: str) -> bytes:
    fig = render_onepage(o)
    buf = io.BytesIO()
    fig.savefig(buf, format=fmt, dpi=DPI)
    plt.close(fig)
    return buf.getvalue()


def para_png(o: dict) -> bytes:
    return _salvar(o, "png")


def para_pdf(o: dict) -> bytes:
    return _salvar(o, "pdf")


def para_pptx(o: dict, png: bytes | None = None) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches
    png = png or para_png(o)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    # A4 (1,414) num slide 16:9 (1,778): ajusta pela altura e centraliza
    h = prs.slide_height
    w = int(h * A4[0] / A4[1])
    s.shapes.add_picture(io.BytesIO(png), int((prs.slide_width - w) / 2), 0, width=w, height=h)
    s.notes_slide.notes_text_frame.text = "\n".join(sem_markdown(f) for f in o["frases"])
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def nome_arquivo(o: dict, ext: str) -> str:
    return f"OnePage_Itirapina_Pederneiras_{o['data_atual']:%Y-%m-%d}.{ext}"


def textos_fora_do_quadro(fig) -> list[str]:
    """Textos cujo bbox sai da figura (teste de aceitação)."""
    fig.canvas.draw()
    r = fig.canvas.get_renderer()
    W, H = fig.bbox.width, fig.bbox.height
    fora = []
    textos = list(fig.texts) + [t for ax in fig.axes for t in ax.texts]
    for t in textos:
        if not t.get_text().strip():
            continue
        b = t.get_window_extent(renderer=r)
        if b.x0 < -1 or b.y0 < -1 or b.x1 > W + 1 or b.y1 > H + 1:
            fora.append(t.get_text()[:40])
    return fora
